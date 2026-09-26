from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

F = "Meiryo UI"
navy = PatternFill("solid", fgColor="1F3864")
lav = PatternFill("solid", fgColor="D9D2E9")
gold = PatternFill("solid", fgColor="FFF2CC")
gray = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="A6A6A6")
bd = Border(left=thin, right=thin, top=thin, bottom=thin)
ctr = Alignment(horizontal="center", vertical="center", wrap_text=True)
lft = Alignment(horizontal="left", vertical="center", wrap_text=True)


def header(ws, row, vals):
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font = Font(name=F, bold=True, color="FFFFFF", size=11)
        c.fill = navy
        c.alignment = ctr
        c.border = bd


def cell(ws, r, col, v, align=ctr, fill=None, bold=False, size=10):
    c = ws.cell(row=r, column=col, value=v)
    c.font = Font(name=F, size=size, bold=bold)
    c.alignment = align
    c.border = bd
    if fill:
        c.fill = fill
    return c


wb = Workbook()

# ---------- Sheet 1: 週間フロー ----------
ws = wb.active
ws.title = "週間フロー"
ws["A1"] = "宇宙銀行 週間フロー（2026/9/27改訂）"
ws["A1"].font = Font(name=F, bold=True, size=14)
ws["A2"] = "紫＝作業・公開がある枠／黄＝今回の変更点"
ws["A2"].font = Font(name=F, size=9, color="595959")
header(ws, 3, ["種別", "月", "火", "水", "木", "金", "土", "日"])

L, Y, N = lav, gold, None
rows = [
    ("3時間動画", [
        ("－", N), ("－", N), ("制作（木曜分）", L),
        ("投稿（臨時収入）\n制作（金曜分）", L), ("投稿（高額当選）", L),
        ("制作（日曜分）", L), ("投稿（不安解消）", L)]),
    ("成功神殿", [
        ("投稿（出発・今週の流れ）\n6:00", L), ("制作（木曜分）", L), ("－", N),
        ("投稿（継続・確かな自信）\n6:00", L), ("制作（月曜分）", L), ("－", N), ("－", N)]),
    ("ライブ", [
        ("実施 18:00〜翌5:00\n（11時間・888黄金の扉）\n＋土曜ライブ素材づくり", Y),
        ("5:00終了\n（月曜ライブ）", L),
        ("実施 12:00〜23:00\n（11時間・窓口の再放送）", Y),
        ("－", N), ("－", N),
        ("実施 18:00〜翌5:00\n（11時間・宇宙銀行の窓口）", Y),
        ("5:00終了（土曜ライブ）\n素材準備（月曜分）\nライブ予約", L)]),
    ("コミュニティ投稿", [
        ("直前リマインド\n（17:30）", L),
        ("アンケート調査\n（15:30）", L),
        ("ライブ直前リマインド\n（11:30）＋成功神殿告知", L),
        ("－", N),
        ("アナリティクス分析", L),
        ("窓口リマインド（17:30）\n紙とペンを用意", Y),
        ("週間まとめ＋記帳の紹介\n＋来週予告＋成功神殿告知\n（17:00）", Y)]),
    ("Nightbot", [
        ("共通タイマーON", Y), ("－", N), ("共通タイマーON", Y),
        ("－", N), ("－", N), ("共通タイマーON", Y), ("－", N)]),
    ("ショート動画", [
        ("今週の問いかけ\nCanva制作\n→翌朝7:00予約", Y),
        ("投稿 7:00\n（今週の問いかけ）", Y),
        ("－", N), ("－", N),
        ("窓口予告ショート制作\n→翌朝7:00予約\n（任意・同素材）", Y),
        ("投稿 7:00\n（窓口予告・任意）", Y),
        ("－", N)]),
    ("Xのポスト", [
        ("火曜分を予約\n（ショートと同じ問いかけ）", Y),
        ("投稿 7:00\n（今週の問いかけ）", Y),
        ("－", N), ("－", N), ("－", N),
        ("投稿 12:00\n（窓口リマインド）", Y),
        ("投稿\n（記帳の紹介・任意）", Y)]),
    ("コメント", [("5分いいね！", N)] * 7),
]
for r, (name, cells) in enumerate(rows, 4):
    cell(ws, r, 1, name, fill=gray, bold=True)
    for j, (v, fill) in enumerate(cells, 2):
        cell(ws, r, j, v, fill=fill)
    ws.row_dimensions[r].height = 62
ws.row_dimensions[3].height = 24
ws.column_dimensions["A"].width = 18
for col in "BCDEFGH":
    ws.column_dimensions[col].width = 23

notes = [
    "【変更点】",
    "・9時間動画の行を削除（9/26から土曜はライブに変更）。月曜の「制作（前倒し）」枠は土曜ライブの素材づくりに充てる。ワーク用テロップは初回に作れば毎週使い回し可",
    "・アンケート調査は火曜15:30",
    "・土曜：「9時間動画告知」→「窓口リマインド（紙とペンを用意）」に変更",
    "・水曜ライブは「土曜の窓口」の再放送。土曜に来られなかった人も参加でき、素材も使い回せる",
    "・日曜の週間まとめに「今週の窓口の記帳から」2〜3件の紹介を追加。チャットを読み返すのはこの時だけでOK",
    "・Nightbotは配信中にしか動かないため、共通タイマーは常時ONのままで切り替え作業なし",
    "・ライブはすべて11時間（アーカイブが残る長さ）。月・土は18:00〜翌5:00、水は12:00〜23:00",
    "・白蛇さんは「黄金の扉の番人」として月曜ライブの画像に時々登場させる案（未確定）",
    "・ショートとXを別の行に分けた。どちらも「窓口ワークの問いかけを1つ置いて、土曜の窓口に誘う入口」。文面は「ショート・X」シートの6週ローテーションから選ぶだけ",
]
start = 4 + len(rows) + 1
for i, t in enumerate(notes):
    r = start + i
    c = ws.cell(row=r, column=1, value=t)
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=8)
    c.font = Font(name=F, size=10, bold=(i == 0))
    c.alignment = lft
    ws.row_dimensions[r].height = 20
ws.freeze_panes = "B4"

# ---------- Sheet 2: 窓口ワーク ----------
w2 = wb.create_sheet("窓口ワーク")
w2["A1"] = "宇宙銀行の窓口（土曜18時〜／水曜12時〜再放送）"
w2["A1"].font = Font(name=F, bold=True, size=14)
w2["A2"] = "3時間ブロックに焼き込み、ループのたびに窓口が開く設計（11時間の配信なら、0:00・3:00・6:00・9:00開始の約4回）"
w2["A2"].font = Font(name=F, size=9, color="595959")
header(w2, 3, ["ブロック内の時間", "ワーク", "その場でやること", "チャットへの一言", "動かすビリーフ"])
work = [
    ("0:00〜", "開店の案内", "紙とペンを用意してもらう", "－", "－"),
    ("0:10〜", "財布の窓口整理",
     "レシートを出す・お札の向きを揃える・カードを1枚減らす（→5分の静かな作業画面）",
     "「整えました」", "－"),
    ("0:30〜", "記帳ノート",
     "今週受け取った小さなもの（お釣り・ポイント・おすそ分け・褒め言葉）を3つ書き、財布か通帳に挟む",
     "受け取ったものを1つ", "木「お金は苦労しないと入ってこない」"),
    ("0:40〜", "今週のラッキー",
     "今週あった小さなラッキーを1つ思い出す／当選したら最初にすることを1つ書く",
     "ラッキーを1つ", "金「自分には運がない」"),
    ("0:50〜", "残高確認",
     "アプリか通帳を開く。①怖さを認める ②数字はあなたの価値ではない ③開く前に3回息を吐き「ありがとう」 ④見たら「ここから満ちていく」",
     "「見ました」\n※金額は書かない", "日「お金のことを考えるのは怖い」\n土「お金はいつか尽きる」"),
    ("1:10〜", "不安のお預け入れ",
     "お金の不安を1つ紙に書き、折りたたんで封筒か引き出しへ。「宇宙銀行にお預けしました」",
     "「預けました」", "日「お金のことを考えるのは怖い」"),
    ("1:20〜", "来週の宿題",
     "次にお金を払うとき、心の中で「ありがとう」と言う。来週の窓口で報告",
     "「できました」（翌週）", "－"),
    ("1:30〜3:00", "閉店・おやすみモード", "通常のBGMループ", "－", "－"),
]
for r, row in enumerate(work, 4):
    for j, v in enumerate(row, 1):
        cell(w2, r, j, v, align=ctr if j in (1, 2) else lft)
    w2.row_dimensions[r].height = 58
for col, wd in zip("ABCDE", [16, 18, 60, 24, 34]):
    w2.column_dimensions[col].width = wd

r0 = 13
w2.cell(row=r0, column=1, value="【曜日テーマとビリーフの対応】").font = Font(name=F, bold=True)
header(w2, r0 + 1, ["曜日", "今のテーマ", "動かしたいビリーフ", "対応する窓口ワーク"])
bel = [
    ("木", "臨時収入", "「お金は苦労しないと入ってこない」", "記帳ノート"),
    ("金", "高額当選", "「自分には運がない」", "今週のラッキー"),
    ("土", "一生お金に困らない安定", "「お金はいつか尽きる・足りなくなる」", "残高確認（すでにあるもの）"),
    ("日", "不安解消", "「お金のことを考えるのは怖い」", "残高確認／不安のお預け入れ"),
]
for i, row in enumerate(bel, r0 + 2):
    for j, v in enumerate(row, 1):
        cell(w2, i, j, v, align=ctr if j < 3 else lft)
    w2.row_dimensions[i].height = 24

# ---------- Sheet 3: Nightbot設定 ----------
w3 = wb.create_sheet("Nightbot設定")
w3["A1"] = "Nightbot設定メモ"
w3["A1"].font = Font(name=F, bold=True, size=14)
info = [
    "・YouTubeでは「配信中」かつ「公開」のライブにだけ参加する（限定公開のテストでは動かない）",
    "・使う前に、チャンネルでNightbotをモデレーターにする",
    "・タイマーは、指定時間が経ったあと直近5分のチャット数が設定値に達しているか確認してから投稿する。夜中は静かなので「チャット行数」は最小値に",
]
for i, t in enumerate(info, 2):
    c = w3.cell(row=i, column=1, value=t)
    w3.merge_cells(start_row=i, start_column=1, end_row=i, end_column=3)
    c.font = Font(name=F, size=10)
    c.alignment = lft
    w3.row_dimensions[i].height = 30

w3.cell(row=6, column=1, value="① 共通タイマー（常時ON）").font = Font(name=F, bold=True)
header(w3, 7, ["名前", "間隔", "メッセージ"])
tm = [
    ("welcome", "30〜60分", "✨宇宙銀行へようこそ。今夜は流したまま、おやすみくださいね"),
    ("kicho", "30〜60分", "📖今週受け取った小さなものを、ひとつ書いてみませんか。※金額など個人情報は書かないでください"),
    ("madoguchi", "30〜60分", "🏦土曜18時・水曜12時は『宇宙銀行の窓口』。紙とペンを用意してお越しください"),
]
for i, row in enumerate(tm, 8):
    for j, v in enumerate(row, 1):
        cell(w3, i, j, v, align=ctr if j < 3 else lft)
    w3.row_dimensions[i].height = 34

w3.cell(row=12, column=1, value="② コマンド（視聴者が打つと自動返信）").font = Font(name=F, bold=True)
header(w3, 13, ["コマンド", "用途", "返信メッセージ"])
cm = [
    ("!madoguchi", "今日の窓口ワーク",
     "🏦窓口ワーク：①財布を整える ②受け取ったものを3つ書く ③残高を見る（開く前に3回息を吐いて「ありがとう」） ④不安を紙に書いて預ける"),
    ("!yokin", "不安のお預け入れ",
     "🌙お金の不安をひとつ紙に書いて、折りたたんでしまってください。「宇宙銀行にお預けしました」で完了です"),
    ("!uketorimasu", "受け取り", "✨受け取りました。あなたの申込書は宇宙銀行に届いています"),
]
for i, row in enumerate(cm, 14):
    for j, v in enumerate(row, 1):
        cell(w3, i, j, v, align=ctr if j < 3 else lft)
    w3.row_dimensions[i].height = 44
for col, wd in zip("ABC", [18, 20, 80]):
    w3.column_dimensions[col].width = wd

# ---------- Sheet 4: 視聴者の体験 ----------
w4 = wb.create_sheet("視聴者の体験", 1)
w4["A1"] = "視聴者さんが窓口でやること・得られるもの"
w4["A1"].font = Font(name=F, bold=True, size=14)
w4["A2"] = "所要時間は合計15分ほど。用意するのは紙とペン、お財布、通帳かアプリだけ"
w4["A2"].font = Font(name=F, size=9, color="595959")
header(w4, 3, ["ワーク", "具体的にやること", "得られるもの", "ビフォー", "アフター"])
exp = [
    ("財布の窓口整理",
     "レシートを全部出す／お札の向きを揃える／使わないカードを1枚抜く",
     "お金を「雑に扱っていた自分」に気づく。整えた達成感がすぐ手に入る",
     "財布がパンパンで、中身を把握していない",
     "開くたびに気持ちいい財布。お金を丁寧に扱う感覚が戻る"),
    ("記帳ノート",
     "今週受け取った小さなもの（お釣り・ポイント・おすそ分け・褒め言葉）を3つ紙に書き、財布に挟む",
     "「入ってきているもの」が見えるようになる。臨時収入に気づくアンテナが立つ",
     "「お金は出ていくばかり」「何も入ってこない」",
     "「意外と受け取ってた」。小さな入金に日常で気づけるようになる"),
    ("今週のラッキー",
     "小さなラッキーを1つ思い出して書く／当選したら最初にすることを1つ決める",
     "自分は運がいい側の人間だという実感。受け取る準備ができる",
     "「どうせ私には当たらない」",
     "「私にもいいこと起きてる」。受け取った後の自分がイメージできる"),
    ("残高確認",
     "開く前に3回息を吐いて「ありがとう」→ アプリか通帳を見る →「ここから満ちていく」と声をかける",
     "数字と自分の価値を切り離せる。見ても大丈夫だった、という成功体験",
     "残高を見るのが怖くて、アプリを開けない",
     "週1回、落ち着いて残高を確認できる自分になる"),
    ("不安のお預け入れ",
     "お金の不安を1つ紙に書き、折りたたんで封筒や引き出しにしまう",
     "頭の中でぐるぐるしていた不安が「外に出る」。寝つきがよくなる",
     "寝る前にお金の心配が止まらない",
     "「預けたから大丈夫」と手放して眠れる"),
    ("来週の宿題",
     "次にお金を払うとき、心の中で「ありがとう」と言う。翌週チャットで報告",
     "支払いが「減る」から「めぐる」に変わる。毎週窓口に来る理由ができる",
     "支払いのたびに、ちょっと胸がざわつく",
     "支払いのたびに感謝できる。お金との関係がやわらかくなる"),
]
for r, row in enumerate(exp, 4):
    for j, v in enumerate(row, 1):
        cell(w4, r, j, v, align=ctr if j == 1 else lft, fill=gray if j == 1 else None, bold=(j == 1))
    w4.row_dimensions[r].height = 60

r0 = 11
w4.cell(row=r0, column=1, value="【通い続けたときの変化（全体のビフォーアフター）】").font = Font(name=F, bold=True)
header(w4, r0 + 1, ["", "ビフォー（初めて来た日）", "1か月後", "3か月後", ""])
arc = [
    ("お金への気持ち", "漠然と不安。考えたくない", "小さな入金に気づける", "お金の話を落ち着いて考えられる"),
    ("行動", "BGMを流すだけ", "毎週窓口で手を動かす", "残高確認・財布整理が習慣になる"),
    ("自分への見方", "「私はお金に縁がない」", "「意外と受け取ってた」", "「私は受け取っていい人」"),
    ("チャンネルとの関係", "たまたま流れてきた視聴者", "土曜が楽しみな常連", "記帳を報告し合う仲間"),
]
for i, row in enumerate(arc, r0 + 2):
    for j, v in enumerate(row, 1):
        cell(w4, i, j, v, align=ctr if j == 1 else lft, fill=gray if j == 1 else None, bold=(j == 1))
    w4.row_dimensions[i].height = 30
w4.cell(row=r0 + 7, column=1,
        value="ひとことで：「BGMを流して待つ人」から「お金と向き合い、気づきを受け取れる人」へ").font = Font(name=F, bold=True, size=11)
for col, wd in zip("ABCDE", [18, 40, 36, 32, 36]):
    w4.column_dimensions[col].width = wd

# ---------- Sheet: ショート・X ----------
w6 = wb.create_sheet("ショート・X", 2)
w6["A1"] = "ショート・X 運用フロー（窓口への入口）"
w6["A1"].font = Font(name=F, bold=True, size=14)
w6["A2"] = "役割：ワークを全部やってもらう場所ではなく「問いかけを1つ置いて、土曜の窓口に誘う」入口。顔出し・声なし"
w6["A2"].font = Font(name=F, size=9, color="595959")

w6.cell(row=4, column=1, value="【毎週のフロー】").font = Font(name=F, bold=True, size=11)
header(w6, 5, ["曜日", "ステップ", "やること", "所要時間"])
flow = [
    ("月", "① 今週のワークを選ぶ", "下の6週ローテーション表から、今週の番号の行を開く", "1分"),
    ("月", "② ショートを作る", "Canvaのテンプレで、テロップ1〜3と締めカードの文字だけ差し替え。BGMは既存音源を15〜30秒切り出す", "15分"),
    ("月", "③ 予約する", "ショートを火曜7:00に予約。Xの問いかけ文（下表）をコピペして火曜7:00に予約", "5分"),
    ("金", "④ 窓口予告（任意）", "同じショートの締めカードを「明日18時、窓口が開きます」に差し替えて土曜7:00に予約", "5分"),
    ("土", "⑤ Xリマインド", "12:00に「今夜18時、宇宙銀行の窓口が開きます。紙とペンを用意してお越しください🏦」を投稿", "1分"),
    ("日", "⑥ 記帳の紹介（任意）", "窓口のチャットから素敵な書き込みを1つ選び、名前を伏せてXで紹介", "5分"),
]
for r, row in enumerate(flow, 6):
    for j, v in enumerate(row, 1):
        cell(w6, r, j, v, align=ctr if j in (1, 4) else lft)
    w6.row_dimensions[r].height = 36

r0 = 13
w6.cell(row=r0, column=1, value="【6週ローテーション】1週に1ワーク。6週で一周したら第1週に戻る（同じ問いかけの繰り返しでOK）").font = Font(name=F, bold=True, size=11)
header(w6, r0 + 1, ["週", "ワーク", "ショート テロップ1（問いかけ）", "ショート テロップ2（例）", "ショート テロップ3（行動）", "X投稿文（140字以内）"])
rot = [
    (1, "記帳ノート",
     "今週、お金の形をしていない\n「入金」はありましたか？",
     "お釣りの端数\nおすそ分け\n褒め言葉",
     "受け取ったものを\nひとつ書いてみてください",
     "「お金は苦労しないと入ってこない」そう思っていませんか。\n今週、頑張らずに受け取ったものをひとつ思い出してみてください。お釣りの端数、おすそ分け、褒め言葉も立派な入金です。\n土曜18時、宇宙銀行の窓口で🏦\n#宇宙銀行 #金運"),
    (2, "今週のラッキー",
     "「どうせ私には当たらない」\nと思っていませんか？",
     "信号がすぐ青になった\n欲しかったものが安かった",
     "今週の小さなラッキーを\nひとつ思い出してください",
     "「どうせ私には当たらない」と思う前に。\n今週あった小さなラッキーをひとつ思い出してみてください。気づいた分だけ、運はあなたの味方になります。\n土曜18時、宇宙銀行の窓口で🏦\n#宇宙銀行 #金運"),
    (3, "残高確認",
     "残高を見るのが\n怖くありませんか？",
     "それは、お金を\n大切にしてきた証です",
     "開く前に3回息を吐いて\n「ありがとう」",
     "残高を見るのが怖いのは、お金を大切にしてきた証です。\n開く前に3回息を吐いて「ありがとう」。見たら「ここから満ちていく」。\n数字はあなたの価値ではありません。\n土曜18時、窓口で🏦\n#宇宙銀行 #金運"),
    (4, "不安のお預け入れ",
     "寝る前に\nお金の心配が止まらない夜は",
     "不安をひとつ紙に書いて\n折りたたんでください",
     "「宇宙銀行に\nお預けしました」",
     "寝る前にお金の心配が止まらない夜は、不安をひとつ紙に書いて、折りたたんでしまってください。\n「宇宙銀行にお預けしました」\n消さなくていい、預けるだけ。\n土曜18時、窓口で🏦\n#宇宙銀行 #金運"),
    (5, "財布の窓口整理",
     "お財布、パンパンに\nなっていませんか？",
     "レシートを出す\nお札の向きを揃える\nカードを1枚抜く",
     "5分で、お金を丁寧に\n扱う感覚が戻ります",
     "財布、パンパンになっていませんか。\nレシートを出す。お札の向きを揃える。カードを1枚抜く。\n5分で、お金を丁寧に扱う感覚が戻ってきます。\n土曜18時、宇宙銀行の窓口で🏦\n#宇宙銀行 #金運"),
    (6, "ありがとう払い",
     "今週の宿題です",
     "次にお金を払うとき\n心の中で「ありがとう」",
     "支払いは減るのではなく\nめぐっていくもの",
     "今週の宿題です。\n次にお金を払うとき、心の中で「ありがとう」と言ってみてください。\n支払いは減るのではなく、めぐっていくもの。\n土曜18時、宇宙銀行の窓口で🏦\n#宇宙銀行 #金運"),
]
for r, row in enumerate(rot, r0 + 2):
    for j, v in enumerate(row, 1):
        cell(w6, r, j, v, align=ctr if j in (1, 2) else lft,
             fill=gray if j in (1, 2) else None, bold=(j == 2))
    w6.row_dimensions[r].height = 110
rr = r0 + 2 + len(rot) + 1
w6.cell(row=rr, column=1, value="【締めカード（全週共通）】").font = Font(name=F, bold=True)
w6.cell(row=rr + 1, column=1, value="火曜ショート：「続きは土曜18時 宇宙銀行の窓口で」／土曜予告ショート：「今夜18時、窓口が開きます。紙とペンを用意してお越しください」").font = Font(name=F, size=10)
w6.merge_cells(start_row=rr + 1, start_column=1, end_row=rr + 1, end_column=6)
for col, wd in zip("ABCDEF", [8, 18, 28, 28, 28, 60]):
    w6.column_dimensions[col].width = wd

# ---------- Sheet: 収益化 ----------
w7 = wb.create_sheet("収益化")
w7["A1"] = "収益化の仕組み（張り付かなくても回る形）"
w7["A1"].font = Font(name=F, bold=True, size=14)
w7["A2"] = "大前提：「スパチャすると金運が上がる」とは言わない。あくまで「応援・感謝」として受け取る（開運商法に見えない／YouTubeの誤解を招くコンテンツ規約対策）"
w7["A2"].font = Font(name=F, size=10, bold=True, color="C00000")
w7.merge_cells("A2:E2")
w7["A2"].alignment = lft
w7.row_dimensions[2].height = 30
header(w7, 4, ["優先", "仕組み", "中身", "視聴者が得るもの", "手間"])
mon = [
    ("1", "メンバーシップ\n「窓口会員」",
     "・会員スタンプ（金貨・宇宙銀子・扉・宝石のカスタム絵文字）\n・会員限定：窓口ワークシートPDF／月1回の会員限定窓口ライブ（素材使い回し）／限定長尺版\n・バッジ名：普通預金 → 定期預金 → 金庫会員（継続月数で変化）",
     "スタンプで参加している実感、限定ワーク、続けるほど上がる「通帳ランク」",
     "初回の設定とスタンプ作成のみ。以降は月1回の限定ライブ予約"),
    ("2", "ワーク帳の商品化",
     "・窓口ワークをまとめて Kindle本 or 印刷用ノートPDF（BOOTH／Gumroad）に\n・概要欄・固定コメント・会員特典からリンク",
     "ライブ以外の日も一人でワークできる手元のノート",
     "作って置くだけ。Kindle出版のノウハウをそのまま活用"),
    ("3", "スパチャ\n「宇宙銀行へのお預け入れ」",
     "・OBSにアラート（Streamlabs等）を入れ、金貨演出＋「◯◯様のお預け入れを承りました✨」を自動表示\n・その週の方のお名前を翌週の窓口ライブ背景に「預金者名簿」としてエンドロール表示\n・金額帯の呼び名：「感謝の一滴」「満ちる預け入れ」など世界観に合わせる",
     "画面に名前が出る喜び、名簿に刻まれる特別感",
     "初回設定のみ自動。名簿は週1回まとめて作成"),
]
for r, row in enumerate(mon, 5):
    for j, v in enumerate(row, 1):
        cell(w7, r, j, v, align=ctr if j in (1, 2) else lft,
             fill=gray if j in (1, 2) else None, bold=(j in (1, 2)))
    w7.row_dimensions[r].height = 100
w7.cell(row=9, column=1, value="※今のライブ規模（同時視聴100人前後）ではスパチャ単体はまだ小さめ。まずは毎月の土台になるメンバーシップから。").font = Font(name=F, size=10)
w7.merge_cells("A9:E9")
for col, wd in zip("ABCDE", [7, 22, 70, 36, 34]):
    w7.column_dimensions[col].width = wd

# ---------- Sheet: やることリスト ----------
w8 = wb.create_sheet("やることリスト", 1)
w8["A1"] = "やることリスト"
w8["A1"].font = Font(name=F, bold=True, size=14)
w8["A2"] = "終わったら「済」欄に ✓ を入れてください"
w8["A2"].font = Font(name=F, size=9, color="595959")
header(w8, 3, ["済", "区分", "やること", "メモ"])
todo = [
    ("土曜窓口ライブ", "窓口ワークのテロップを作る（案内3〜4枚×6ワーク）", "初回だけ。以降は毎週使い回し"),
    ("土曜窓口ライブ", "3時間ブロックにワークを焼き込んだ窓口用素材を書き出す", "「窓口ワーク」シートの時間割どおり"),
    ("土曜窓口ライブ", "土曜18時のライブを予約（自動スタートON）", ""),
    ("土曜窓口ライブ", "固定コメント「金額など個人情報は書かないでください」を用意", ""),
    ("水曜ライブ", "水曜12時を「窓口の再放送」として予約", "土曜と同じ素材"),
    ("Nightbot", "nightbot.tv にYouTubeでログインし、Nightbotをモデレーターにする", ""),
    ("Nightbot", "共通タイマー3つを登録（チャット行数は最小値）", "「Nightbot設定」シート①"),
    ("Nightbot", "コマンド3つを登録（!madoguchi / !yokin / !uketorimasu）", "「Nightbot設定」シート②"),
    ("Nightbot", "公開ライブ中に動くか一度確認する", "限定公開では動かない"),
    ("ショート・X", "Canvaでショートのテンプレートを1つ作る", "問いかけ・例・行動・締めカードの4枚"),
    ("ショート・X", "第1週（記帳ノート）のショートとXを予約", "「ショート・X」シート"),
    ("収益化", "メンバーシップの利用条件を確認し、有効にする", "YouTube Studio → 収益化 → メンバーシップ"),
    ("収益化", "会員スタンプ（カスタム絵文字）を作る", "金貨・宇宙銀子・扉・宝石など"),
    ("収益化", "会員バッジの名前を設定（普通預金 → 定期預金 → 金庫会員）", ""),
    ("収益化", "OBSにスパチャのアラート演出を入れる", "Streamlabs など"),
    ("収益化", "「預金者名簿」のエンドロール用テンプレートを作る", "週1回まとめて更新"),
    ("収益化", "窓口ワーク帳の構成を考える", "窓口ワーク6つ＋書き込みページ"),
    ("収益化", "ワーク帳をAmazonでペーパーバック＋Kindle出版", "「販売先・メンバーシップ」シート"),
    ("収益化", "Content IDの申し立てがある動画を確認（メンバーシップの対象外になるため）", "YouTube Studio → コンテンツ → 制限"),
    ("告知", "コミュニティ投稿で土曜窓口の開始を告知（紙とペンを用意）", "初回は金曜夜に"),
    ("未決定", "白蛇さんの位置づけを決める", "案：黄金の扉の番人"),
]
for r, (cat, task, memo) in enumerate(todo, 4):
    cell(w8, r, 1, "□")
    cell(w8, r, 2, cat, fill=gray)
    cell(w8, r, 3, task, align=lft)
    cell(w8, r, 4, memo, align=lft)
    w8.row_dimensions[r].height = 24
for col, wd in zip("ABCD", [6, 16, 62, 40]):
    w8.column_dimensions[col].width = wd
w8.freeze_panes = "A4"

# ---------- Sheet 5: 参考チャンネル10選 ----------
w5 = wb.create_sheet("参考チャンネル10選")
w5["A1"] = "似たことをしているチャンネル10選"
w5["A1"].font = Font(name=F, bold=True, size=14)
w5["A2"] = "「BGM×張り付かないワーク」そのものは見当たらず、要素ごとに近いものを選定。登録者数・ライブ頻度は未確認（2026/9/27調べ）"
w5["A2"].font = Font(name=F, size=9, color="595959")
header(w5, 3, ["No.", "チャンネル", "区分", "どんな内容か", "宇宙銀行で参考にできる点", "参考URL"])
ch = [
    (1, "心屋仁之助【公式】", "お金のブロック（日本）",
     "「お金に恵まれない」「口ぐせを変えたい」などの公開カウンセリング",
     "ビリーフを動かす言葉づかい・問いかけ方",
     "https://www.youtube.com/@kokoro-ya"),
    (2, "本田晃一", "お金のブロック（日本）",
     "「世界一ゆる〜い幸せの帝王学 お金編」",
     "がんばらずに受け取る語り口（木曜テーマに近い）",
     "https://www.youtube.com/watch?v=BoY3bo3DHBw"),
    (3, "本田健「Dear Ken」", "お金のブロック（日本）",
     "「お金のブロックを外すには？」などの人生相談",
     "視聴者の悩みの言語化の仕方",
     "https://www.youtube.com/watch?v=yswREi_4fNc"),
    (4, "【潜在意識・引き寄せ】探究ちゃんねる", "潜在意識（日本）",
     "潜在意識の書き換えを解説",
     "潜在意識の説明の仕方（概要欄・テロップ）",
     "https://www.youtube.com/channel/UCsDfqx5o9Z4Qac_UOGxEfLw"),
    (5, "開運＆潜在意識チャンネル", "潜在意識（日本）",
     "開運と潜在意識の掛け合わせ",
     "テーマが最も近い。タイトル・サムネの比較対象に",
     "https://www.youtube.com/channel/UCMlFv1Ou-MnfDs7w-ETX1cQ"),
    (6, "たかみー💰金運上昇チャンネル", "金運実践（日本）",
     "財布・金運の実践系（※Xのプロフィールから確認。YouTube側の中身は未確認）",
     "財布ワークの見せ方",
     "https://x.com/UP_kinun"),
    (7, "Denise Duffield-Thomas", "お金のブロック（海外）",
     "money blocks（お金のブロック）解除の第一人者",
     "よくあるブロック7つの整理 → 曜日別ビリーフの補強に",
     "https://www.youtube.com/channel/UCgu5fJfr91Q0rjewdtlxUtg"),
    (8, "Kate Northrup", "お金と向き合う習慣（海外）",
     "著書『Money: A Love Story』。定期的にお金と向き合う「お金とのデート」を提唱",
     "「窓口」の考え方に最も近い。週1で通う習慣づくり",
     "https://www.youtube.com/watch?v=qFafymotfSA"),
    (9, "MONETS MONEY", "一緒に手を動かす（海外）",
     "「Budget With Me」。現金を封筒に分ける様子を毎回シリーズ配信",
     "「見ながら一緒にやる」形式のシリーズ化",
     "https://www.youtube.com/watch?v=dG6oE7rJjQg"),
    (10, "Jules Acree", "一緒に手を動かす（海外）",
     "「Real Time Journal with Me」。問いかけを出してその場で一緒に書く",
     "台本を焼き込む設計に最も近い。問いかけ→書く時間の間合い",
     "https://www.julesacree.com/all/journal-with-me"),
]
for r, row in enumerate(ch, 4):
    for j, v in enumerate(row, 1):
        c = cell(w5, r, j, v, align=ctr if j in (1, 3) else lft)
        if j == 6:
            c.hyperlink = v
            c.font = Font(name=F, size=9, color="0563C1", underline="single")
    w5.row_dimensions[r].height = 44
for col, wd in zip("ABCDEF", [6, 28, 20, 44, 40, 42]):
    w5.column_dimensions[col].width = wd
w5.cell(row=15, column=1,
        value="【ヒント】海外定番の「〜 with me」形式は、配信者が話さなくても視聴者が参加できることを示している。BGMと組み合わせたチャンネルは見当たらず、宇宙銀行の窓口は空いているポジション").font = Font(name=F, bold=True, size=10)
w5.merge_cells(start_row=15, start_column=1, end_row=15, end_column=6)
w5.cell(row=15, column=1).alignment = lft
w5.row_dimensions[15].height = 34

# ---------- Sheet: 販売先・メンバーシップ ----------
w9 = wb.create_sheet("販売先・メンバーシップ")
w9["A1"] = "ワーク帳の販売先と、メンバーシップの設計"
w9["A1"].font = Font(name=F, bold=True, size=14)
w9["A2"] = "視聴者の中心は55〜64歳・65歳以上。「新しいアカウントを作らなくていい」「紙で書ける」を優先して選ぶ（手数料は2026年9月時点の調べ）"
w9["A2"].font = Font(name=F, size=9, color="595959")
header(w9, 4, ["おすすめ", "販売先", "どんなところ", "手数料の目安", "チャンネル世代との相性"])
plat = [
    ("◎", "Amazon KDP\n（Kindle＋ペーパーバック）",
     "Amazonで電子書籍と紙の本を販売。紙の本は注文が入るたびに印刷されるので在庫なし",
     "印刷代を引いた残りから、Amazonのロイヤリティ条件に沿って受け取り（Kindle出版の経験がそのまま使える）",
     "◎ ほとんどの人がAmazonのアカウントを持っている。紙のワーク帳に書き込めるので窓口ワークと相性抜群"),
    ("○", "note",
     "日本の文章・コンテンツ販売サービス。有料記事にPDFを添付して販売できる",
     "クレカ決済で手取り約85%（決済手数料5%＋利用料10%）。振込ごとに270円",
     "○ 名前を知っている人が多く、スマホでそのまま買える。PDFのダウンロードに慣れていない人には少し難しい"),
    ("△", "BOOTH",
     "pixiv運営のクリエイター向けショップ。デジタル商品もグッズも販売できる",
     "5.6%＋45円（2025年10月から）",
     "△ 購入にpixivアカウントが必要。イラスト・同人の若い層向けの印象が強い"),
    ("海外向け", "Gumroad",
     "海外向けのデジタル商品販売サービス。英語・ドル建て",
     "10%＋0.50ドル＋カード決済手数料",
     "× 日本の視聴者には不向き。海外向けBGM音源の販売を始めるときの候補"),
]
for r, row in enumerate(plat, 5):
    for j, v in enumerate(row, 1):
        cell(w9, r, j, v, align=ctr if j in (1, 2) else lft,
             fill=gray if j in (1, 2) else None, bold=(j in (1, 2)))
    w9.row_dimensions[r].height = 62
w9.cell(row=9, column=1, value="→ まずは Amazon のペーパーバック（＋Kindle）で「宇宙銀行 窓口ワーク帳」。PDF版は note でメンバー特典や単品販売に回す").font = Font(name=F, bold=True, size=10)
w9.merge_cells("A9:E9")

w9.cell(row=11, column=1, value="【メンバーシップ「窓口会員」の設計】").font = Font(name=F, bold=True, size=12)
header(w9, 12, ["項目", "設定内容", "", "", "メモ"])
mem = [
    ("利用条件", "YouTubeパートナープログラム参加中／18歳以上／子ども向け設定なし",
     "音楽の申し立て（Content ID）がある動画は対象外になるので、申し立て動画は整理しておく"),
    ("レベル数", "最初は1つだけ（月290円前後）",
     "最大6つまで作れるが、はじめはシンプルに。増やすのは会員が増えてから"),
    ("レベル名", "窓口会員", "チャンネルの世界観（銀行）と合わせる"),
    ("特典①", "会員スタンプ（カスタム絵文字）",
     "最初は4個まで。金貨・扉・宝石・「受け取ります」文字など。画像は48×48〜480×480ピクセル"),
    ("特典②", "継続月数で変わる会員バッジ",
     "新規＝口座開設／1か月＝普通預金／2か月＝積立預金／6か月＝定期預金／1年＝金庫会員／2年〜＝宇宙金庫会員"),
    ("特典③", "会員限定のコミュニティ投稿（月1回）",
     "窓口ワークシートPDFや「今月の問いかけ」を配布"),
    ("特典④", "会員限定の窓口ライブ（月1回）",
     "素材は土曜窓口の使い回しでOK"),
    ("お迎えの一言", "「宇宙銀行の窓口会員へようこそ。あなたの口座が開設されました✨」",
     "入会したときの案内文"),
]
for r, (k, v, m) in enumerate(mem, 13):
    cell(w9, r, 1, k, fill=gray, bold=True)
    c = cell(w9, r, 2, v, align=lft)
    w9.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
    cell(w9, r, 5, m, align=lft)
    w9.row_dimensions[r].height = 40
for col, wd in zip("ABCDE", [14, 26, 40, 34, 52]):
    w9.column_dimensions[col].width = wd

for s in wb.worksheets:
    s.page_setup.orientation = "landscape"
    s.page_setup.fitToWidth = 1
    s.sheet_properties.pageSetUpPr.fitToPage = True
    s.sheet_view.showGridLines = False

import os
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), "宇宙銀行_週間フロー.xlsx"))
print("saved")
